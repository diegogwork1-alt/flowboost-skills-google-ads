#!/usr/bin/env python3
"""Calcula la estacionalidad de una cuenta de Google Ads desde la hoja del cliente.

Baja la hoja con rclone (todas las pestañas), lee `datos-google-terminos` y
`datos-google-is`, y saca el índice de estacionalidad mes a mes.

    python3 estacionalidad.py <URL_O_ID_DE_LA_HOJA>
    python3 estacionalidad.py <URL_O_ID> --familias familias.json
    python3 estacionalidad.py <URL_O_ID> --remoto midrive:

`familias.json` es opcional: {"urgencias": ["urgente","averia"], "precio": ["precio","cuanto"]}
Sin él, agrupa por la palabra más significativa de cada término y lo dice.

Requiere: rclone configurado con un remoto de Drive, y openpyxl.
"""
import argparse, json, os, re, subprocess, sys, tempfile
from collections import defaultdict

MESES = ['ene','feb','mar','abr','may','jun','jul','ago','sep','oct','nov','dic']
# Palabras que no definen una familia: aparecen en todo y no dicen nada de la intención
VACIAS = {'de','la','el','en','y','a','para','con','por','que','los','las','un','una','del','al',
          'como','mi','me','se','es','no','si','lo','su','mas','más','o','sin','sobre'}


def resolver_id(txt):
    """Acepta la URL entera de la hoja o solo el id."""
    m = re.search(r'/spreadsheets/d/([A-Za-z0-9_-]{20,})', txt)
    return m.group(1) if m else txt.strip()


def bajar(file_id, remoto, destino):
    cmd = ['rclone', 'backend', 'copyid', remoto, file_id, destino,
           '--drive-export-formats', 'xlsx']
    r = subprocess.run(cmd, capture_output=True, text=True)
    if not os.path.exists(destino) or os.path.getsize(destino) == 0:
        sys.exit(f"✗ rclone no pudo bajar la hoja.\n{r.stderr.strip()[:500]}")
    return destino


def leer(ruta, pestana):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if pestana not in wb.sheetnames:
        return None, wb.sheetnames
    ws = wb[pestana]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return [], wb.sheetnames
    cab = [str(c).strip() if c is not None else '' for c in filas[0]]
    out = [dict(zip(cab, f)) for f in filas[1:] if any(v is not None for v in f)]
    return out, wb.sheetnames


def mes_de(v):
    """La columna `mes` llega como fecha; a veces como texto. Devuelve (año, mes)."""
    if hasattr(v, 'year'):
        return v.year, v.month
    s = str(v)[:10]
    try:
        a, m = s.split('-')[:2]
        return int(a), int(m)
    except Exception:
        return None


def familia_de(termino, familias):
    t = str(termino).lower()
    if familias:
        for nombre, palabras in familias.items():
            if any(p.lower() in t for p in palabras):
                return nombre
        return 'otros'
    # Sin familias definidas: la palabra más larga que no sea vacía ni un número
    palabras = [p for p in re.findall(r'[a-záéíóúñ]+', t) if p not in VACIAS and len(p) > 3]
    return max(palabras, key=len) if palabras else '(sin clasificar)'


def indice(por_mes):
    """índice = (valor del mes ÷ media de todos los meses con dato) × 100"""
    vals = [v for v in por_mes.values() if v is not None]
    if not vals or sum(vals) == 0:
        return {}
    media = sum(vals) / len(vals)
    return {m: round(v / media * 100) for m, v in por_mes.items()} if media else {}


def tabla(titulo, idx, meses_orden):
    print(f"\n  {titulo}")
    print("   " + " ".join(f"{MESES[m-1]:>5}" for m in meses_orden))
    print("   " + " ".join(f"{idx.get(m, '·'):>5}" for m in meses_orden))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('hoja', help='URL o id de la hoja del cliente')
    ap.add_argument('--remoto', default='gdrive:', help='remoto de rclone (por defecto gdrive:)')
    ap.add_argument('--familias', help='JSON con las familias de términos')
    ap.add_argument('--top', type=int, default=12, help='cuántas familias mostrar')
    ap.add_argument('--marca', default='', help='palabras de la marca separadas por comas. '
                    'Los términos que las contengan se analizan APARTE: la marca no es demanda '
                    'de mercado, y mezclarla falsea el índice de todo lo demás')
    args = ap.parse_args()

    familias = json.load(open(args.familias, encoding='utf-8')) if args.familias else None
    marca = [p.strip().lower() for p in args.marca.split(',') if p.strip()]
    file_id = resolver_id(args.hoja)

    tmp = os.path.join(tempfile.mkdtemp(), 'hoja.xlsx')
    print(f"Bajando la hoja {file_id[:12]}… con rclone")
    bajar(file_id, args.remoto, tmp)

    terminos, pestanas = leer(tmp, 'datos-google-terminos')
    if terminos is None:
        sys.exit(f"✗ No existe la pestaña 'datos-google-terminos'.\n"
                 f"  La hoja tiene: {', '.join(pestanas)}\n"
                 f"  → El Ads Script de esa cuenta todavía es el viejo. Actualízalo (ver INSTALACION.md).")
    is_filas, _ = leer(tmp, 'datos-google-is')
    is_filas = is_filas or []

    # ── Agregado ──
    # magnitud → familia → (año,mes) → valor
    agg = {k: defaultdict(lambda: defaultdict(float)) for k in ('impresiones', 'clics', 'conversiones')}
    total = {k: defaultdict(float) for k in ('impresiones', 'clics', 'conversiones', 'coste')}
    meses_vistos = set()

    n_marca = 0
    for f in terminos:
        am = mes_de(f.get('mes'))
        if not am:
            continue
        if marca and any(p in str(f.get('termino')).lower() for p in marca):
            n_marca += 1
            continue          # la marca va aparte: no es demanda del mercado
        meses_vistos.add(am)
        fam = familia_de(f.get('termino'), familias)
        for k in ('impresiones', 'clics', 'conversiones'):
            v = float(f.get(k) or 0)
            agg[k][fam][am] += v
            total[k][am] += v
        total['coste'][am] += float(f.get('coste') or 0)

    if not meses_vistos:
        sys.exit("✗ La pestaña existe pero no tiene datos con fecha. Vuelve a ejecutar el Ads Script.")

    anios = sorted({a for a, _ in meses_vistos})
    n_meses = len(meses_vistos)
    print(f"\n{'═'*66}\nESTACIONALIDAD · {n_meses} meses con datos · años {min(anios)}-{max(anios)}")
    print(f"{len(terminos):,} filas de términos · {len(is_filas):,} filas de cuota de impresiones".replace(',', '.'))
    if marca:
        print(f"{n_marca:,} filas de MARCA excluidas ({', '.join(marca)}): quien te busca por el nombre".replace(',', '.'))
        print("ya te conocía. Eso no es demanda de mercado y falsearía la estacionalidad.")
    else:
        print("\n⚠️  No se ha indicado la marca. Si el cliente recibe búsquedas por su nombre,")
        print("   esas conversiones inflan el índice. Relánzalo con --marca \"nombre,marca\".")

    if n_meses < 24:
        print("\n⚠️  Menos de 24 meses: esto es una HIPÓTESIS, no un patrón confirmado.")
        print("   Con menos de 2 años no se puede separar estacionalidad de tendencia.")
    else:
        print("\n✓ Hay 2 años o más: se puede promediar el mismo mes de varios años.")

    # Promedio del mismo mes across años → índice
    def por_mes_num(d):
        acum = defaultdict(list)
        for (a, m), v in d.items():
            acum[m].append(v)
        return {m: sum(v)/len(v) for m, v in acum.items()}

    meses_orden = sorted({m for _, m in meses_vistos})

    print(f"\n{'─'*66}\nLA CUENTA ENTERA (índice 100 = un mes del montón)")
    for k in ('impresiones', 'clics', 'conversiones'):
        tabla(f"{k:<13} ← {'cuándo BUSCAN' if k=='impresiones' else 'cuándo HACEN CLIC' if k=='clics' else 'cuándo COMPRAN'}",
              indice(por_mes_num(total[k])), meses_orden)

    ic = indice(por_mes_num(total['conversiones']))
    if ic:
        alto = sorted(ic.items(), key=lambda x: -x[1])[:3]
        bajo = sorted(ic.items(), key=lambda x: x[1])[:3]
        print(f"\n  Temporada ALTA (conversiones): {', '.join(f'{MESES[m-1]} ({v})' for m, v in alto)}")
        print(f"  Temporada BAJA (conversiones): {', '.join(f'{MESES[m-1]} ({v})' for m, v in bajo)}")

    # ── Cuota de impresiones perdida por mes ──
    if is_filas:
        pres = defaultdict(list); rank = defaultdict(list)
        for f in is_filas:
            am = mes_de(f.get('mes'))
            if not am:
                continue
            if f.get('is_perdida_presupuesto') is not None:
                pres[am[1]].append(float(f['is_perdida_presupuesto']))
            if f.get('is_perdida_ranking') is not None:
                rank[am[1]].append(float(f['is_perdida_ranking']))
        print(f"\n{'─'*66}\nQUÉ SE PIERDE CADA MES (media de las campañas de Búsqueda)")
        print("   " + " ".join(f"{MESES[m-1]:>5}" for m in meses_orden))
        for nombre, d in (("presupuesto", pres), ("ranking    ", rank)):
            print(f"   " + " ".join(
                f"{round(sum(d[m])/len(d[m])*100):>4}%" if d.get(m) else "    ·" for m in meses_orden))
            print(f"   ↑ perdido por {nombre}")

        print("\n  → Lo perdido por PRESUPUESTO se compra con dinero.")
        print("    Lo perdido por RANKING no se compra: es relevancia, anuncio y landing.")
        if ic:
            fuga = [(m, round(sum(pres[m])/len(pres[m])*100)) for m in meses_orden
                    if pres.get(m) and ic.get(m, 0) >= 100]
            fuga = [x for x in sorted(fuga, key=lambda x: -x[1]) if x[1] >= 10][:4]
            if fuga:
                print("\n  💰 MESES DONDE SE DEJA DINERO (mucha demanda + presupuesto corto):")
                for m, p in fuga:
                    print(f"     {MESES[m-1]}: índice de conversión {ic[m]} · {p}% de impresiones perdidas por presupuesto")

    # ── Familias ──
    print(f"\n{'─'*66}\nPOR FAMILIA DE TÉRMINOS (por conversiones)")
    if not familias:
        print("  (agrupado automáticamente por la palabra principal; para afinar, pasa --familias)")
    ranking = sorted(agg['conversiones'].items(), key=lambda x: -sum(x[1].values()))[:args.top]
    for fam, d in ranking:
        tot = sum(d.values())
        if tot < 1:
            continue
        idx = indice(por_mes_num(d))
        pico = max(idx.items(), key=lambda x: x[1]) if idx else None
        print(f"\n  ▸ {fam}  ({tot:.0f} conversiones en total)"
              + (f" · pico en {MESES[pico[0]-1]} ({pico[1]})" if pico else ""))
        tabla("   índice", idx, meses_orden)

    print(f"\n{'═'*66}")
    print("Los índices son relativos a la propia cuenta: 100 = mes medio, 140 = 40 % por encima.")
    print("Las impresiones dependen del presupuesto que hubo ese mes: léelas con la tabla de arriba.")


if __name__ == '__main__':
    main()
