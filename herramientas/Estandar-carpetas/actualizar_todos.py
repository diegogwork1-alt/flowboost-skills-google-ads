#!/usr/bin/env python3
"""UN SOLO COMANDO que actualiza la pestaña `datos` de TODOS los clientes.

POR QUÉ EXISTE: la tarea programada se quedaba muerta cada vez que iba a ejecutar un
comando que no estaba autorizado (`ls`, `mkdir`, un `python3` distinto…). Con todo el
trabajo detrás de una sola orden hay UN permiso y nada que adivinar.

CÓMO REPARTE LOS DATOS: el agente hace una llamada al MCP de Meta por cliente y cada una
deja su volcado en la misma carpeta, sin decir de quién es. Este script lo averigua solo:
lee las CAMPAÑAS de cada volcado y las compara con las que ya tiene escritas la pestaña
`datos` de cada hoja. Cada volcado va a la hoja donde esas campañas ya viven. Si un volcado
no encaja con ninguna hoja, o encaja con varias, NO se escribe y se dice por qué.

Uso:  actualizar_todos.py [<dir-de-tool-results>] [--dry]
      Sin directorio, usa el de la sesión en curso (el más reciente bajo ~/.claude/projects).
"""
import glob, json, os, re, subprocess, sys, tempfile, time

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
from armar_datos import dec  # noqa: el parseo de números, con sus dos formatos

LISTA = os.path.join(AQUI, "clientes_reportes.json")


def leer_clientes(solo_activos=True):
    """LA lista, en un solo sitio. La comparten este script y la tarea programada; tenerla
    escrita en los dos hacía que se separasen en cuanto entrara un cliente nuevo."""
    d = json.load(open(LISTA))
    return [c for c in d["clientes"]
            if (c.get("activo") and c.get("ad_account_id") and c.get("sheet_id"))
            or not solo_activos]


def dar_de_alta(cliente, cuenta, sheet_id, indicador=""):
    """Suma un cliente a la rutina diaria. Idempotente: repetirlo actualiza, no duplica."""
    d = json.load(open(LISTA))
    cuenta = str(cuenta).replace("act_", "").strip()
    for c in d["clientes"]:
        if c["cliente"].lower() == cliente.lower():
            c.update(ad_account_id=cuenta, sheet_id=sheet_id, activo=True)
            if indicador: c["indicador_leads"] = indicador
            accion = "actualizado"
            break
    else:
        d["clientes"].append({"cliente": cliente, "ad_account_id": cuenta,
                              "sheet_id": sheet_id, "indicador_leads": indicador,
                              "activo": True})
        accion = "dado de alta"
    json.dump(d, open(LISTA, "w"), ensure_ascii=False, indent=2)
    print(f"✓ {cliente} {accion} en la rutina diaria ({len(leer_clientes())} clientes activos)")


def tool_results_por_defecto():
    ds = glob.glob(os.path.expanduser("~/.claude/projects/*/*/tool-results"))
    if not ds: sys.exit("✗ no encuentro ninguna carpeta `tool-results`")
    return max(ds, key=os.path.getmtime)


def campanas_de_la_hoja(sheet_id):
    """Las campañas que YA tiene escritas la pestaña `datos` de esa hoja."""
    import actualizar_datos as ad
    from openpyxl import load_workbook
    tmp = os.path.join(tempfile.mkdtemp(), "h.xlsx")
    ad.bajar(sheet_id, tmp)
    ws = load_workbook(tmp)["datos"]
    col = ad.DATOS.index("campana") + 1
    return {str(ws.cell(r, col).value).strip()
            for r in range(2, ws.max_row + 1) if ws.cell(r, col).value}


def campanas_del_volcado(f):
    try:
        filas = json.loads(json.load(open(f))["ad_entities"])
    except Exception:
        return set()
    return {str(r.get("name", "")).strip() for r in filas
            if dec(r.get("amount_spent")) > 0}


def main():
    if "--listar" in sys.argv:
        for c in leer_clientes():
            print(f'{c["cliente"]}\t{c["ad_account_id"]}\t{c.get("indicador_leads","")}')
        return
    if "--alta" in sys.argv:
        i = sys.argv.index("--alta")
        resto = sys.argv[i+1:]
        if len(resto) < 3:
            sys.exit('uso: --alta "<Cliente>" <ad_account_id> <sheet_id> [indicador_leads]')
        dar_de_alta(resto[0], resto[1], resto[2], resto[3] if len(resto) > 3 else "")
        return
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry" in sys.argv
    dirt = args[0] if args else tool_results_por_defecto()
    volcados = sorted(glob.glob(os.path.join(dirt, "*ads_get_ad_entities*.txt")))
    if not volcados: sys.exit(f"✗ no hay volcados de Meta en {dirt}")
    print(f"carpeta: {dirt}\nvolcados: {len(volcados)}\n")
    CLIENTES = {c["cliente"]: c["sheet_id"] for c in leer_clientes()}

    print("leyendo las campañas de cada hoja…")
    hojas = {}
    for cli, sid in CLIENTES.items():
        try:
            hojas[cli] = campanas_de_la_hoja(sid)
            print(f"  {cli:<16} {len(hojas[cli]):>3} campañas")
        except Exception as e:
            print(f"  {cli:<16} ✗ no se pudo leer: {str(e)[:60]}")

    # a qué cliente pertenece cada volcado
    reparto, sueltos = {}, []
    for v in volcados:
        cam = campanas_del_volcado(v)
        if not cam: continue
        encajan = [c for c, h in hojas.items() if cam & h]
        if len(encajan) == 1:
            reparto.setdefault(encajan[0], []).append(v)
        else:
            sueltos.append((os.path.basename(v), sorted(cam)[:3], encajan))

    print()
    for nom, cam, encajan in sueltos:
        print(f"⚠ volcado sin dueño claro: {nom[-28:]} · campañas {cam} · encaja con {encajan or 'ninguna hoja'}")

    hubo_error = False
    for cli, vs in reparto.items():
        print(f"\n=== {cli} ({len(vs)} volcado/s) ===")
        tmpd = tempfile.mkdtemp()
        for v in vs:                        # aislados: nunca se mezclan dos clientes
            os.symlink(v, os.path.join(tmpd, os.path.basename(v)))
        js = os.path.join(tmpd, "datos.json")
        marca = next((c.get("marca_campanas", "") for c in leer_clientes()
                      if c["cliente"] == cli), "")
        r = subprocess.run([sys.executable, os.path.join(AQUI, "armar_datos.py"),
                            tmpd, js, cli, marca], capture_output=True, text=True)
        print("  " + (r.stdout or r.stderr).strip().replace("\n", "\n  "))
        if r.returncode: hubo_error = True; continue
        cmd = [sys.executable, os.path.join(AQUI, "actualizar_datos.py"),
               CLIENTES[cli], js] + (["--dry"] if dry else [])
        r = subprocess.run(cmd, capture_output=True, text=True)
        print("  " + (r.stdout or r.stderr).strip().replace("\n", "\n  "))
        if r.returncode: hubo_error = True

    faltan = [c for c in CLIENTES if c not in reparto]
    if faltan: print(f"\nsin datos en esta pasada: {', '.join(faltan)}")
    print("\n" + ("terminado CON incidencias" if hubo_error else "terminado sin incidencias"))


if __name__ == "__main__":
    main()
