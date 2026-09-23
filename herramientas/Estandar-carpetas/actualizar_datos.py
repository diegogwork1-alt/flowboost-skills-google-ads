#!/usr/bin/env python3
"""Actualiza la pestaña `datos` de la hoja de un cliente SIN regenerar el libro.

POR QUÉ EXISTE: `montar_hoja_reportes.py` rehace el libro entero, así que sobre una hoja en
uso borra los cierres, los importes a mano y `datos-google`. Esto hace lo que hace n8n cada
mañana —escribir filas en `datos` y corregir las que ya estaban— pero desde aquí, con los
datos que trae el MCP de Meta. No hace falta META_TOKEN.

QUÉ TOCA Y QUÉ NO: escribe **solo** las celdas de la pestaña `datos`. Las pestañas de
reporte son fórmulas que leen de ahí, así que se recalculan solas. `ventas`, `datos-google`
y todo lo escrito a mano se quedan como estaban.

La clave es `id` = «fecha|campaña», igual que en el flujo: un día que ya estaba se CORRIGE
en su sitio, no se duplica. Por eso se puede reescribir la misma ventana todos los días.

Uso:  actualizar_datos.py <sheet_id> <datos.json> [--dry]
"""
import json, os, sys, tempfile, time, urllib.parse, urllib.request
sys.path.insert(0, os.path.expanduser(
    "~/Desktop/FLOWBOOST-BACKUP-MAC/Documentos-Flowboost/Estandar-carpetas"))
from subir_a_drive import api, token, _multipart
from montar_hoja_reportes import DATOS, EUR, ENT, PCT, DEC2, SUAVE, GRIS, NEGRO
from openpyxl import load_workbook

HOJA = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FMTS = [None, None, None, None, EUR, ENT, ENT, DEC2, EUR, ENT, ENT, PCT, ENT, ENT, None]


def bajar(sheet_id, destino):
    u = (f"https://www.googleapis.com/drive/v3/files/{sheet_id}/export?"
         + urllib.parse.urlencode({"mimeType": XLSX}))
    for i in range(6):
        try:
            r = urllib.request.urlopen(urllib.request.Request(
                u, headers={"Authorization": "Bearer " + token(refrescar=(i > 0))}))
            open(destino, "wb").write(r.read()); return
        except Exception as e:
            err = e; time.sleep(1.5 * (i + 1))
    raise RuntimeError(f"no se pudo exportar la hoja: {err}")


def filas_del_json(p):
    """El mismo formato que come `montar_hoja_reportes.py --datos`."""
    d = json.load(open(p)); camps = d["campanas"]
    out = []
    for f in sorted(d["filas"], key=lambda x: (x[0], x[1])):
        nom = camps[f[1]]["nombre"]
        out.append([f"{f[0]}|{nom}", f[0], d["cliente"], nom, f[2], f[3], f[7], f[8],
                    f[9], f[4], f[5], f[10] / 100, f[6], f[11], "MCP " +
                    time.strftime("%Y-%m-%d %H:%M")])
    return out



def _igual(a, b):
    """¿El mismo valor, venga como venga de la hoja?"""
    if a is None: a = 0 if isinstance(b, (int, float)) else ""
    if isinstance(b, (int, float)):
        try: return abs(float(a) - float(b)) < 5e-5
        except (TypeError, ValueError): return False
    return str(a).strip() == str(b).strip()


def _distintas(a, b):
    return any(not _igual(x, y) for x, y in zip(a, b))


def _detalle(antes, fila):
    return "; ".join(f"{DATOS[i]}: {antes[i]}→{fila[i]}"
                     for i in range(len(DATOS) - 1) if not _igual(antes[i], fila[i]))



def _inventario(wb):
    """Qué hay en el libro, para poder comprobar que sigue estando después de tocarlo."""
    inv = {}
    for h in wb.sheetnames:
        ws = wb[h]
        inv[h] = sum(1 for r in ws.iter_rows(min_row=1, max_col=1) if r[0].value is not None)
    return inv


def _comprobar(antes, despues, ws_datos_nuevas):
    """No subir NADA si el ida y vuelta por .xlsx se ha llevado algo por delante.

    El riesgo real de este método: openpyxl reescribe el libro entero, así que un fallo
    suyo podría dejar sin `datos-google` (209 filas en Cliente 03, y solo se recuperan
    reejecutando el script en Google Ads) o sin las pestañas de reporte. Antes de subir se
    comprueba que están TODAS las pestañas y que ninguna ha perdido filas, salvo `datos`,
    que es la única que debe crecer.
    """
    fallos = []
    for h, n in antes.items():
        if h not in despues:
            fallos.append(f"ha desaparecido la pestaña «{h}»")
        elif h != "datos" and despues[h] < n:
            fallos.append(f"«{h}» ha pasado de {n} a {despues[h]} filas")
    if "datos" in antes and despues.get("datos", 0) < antes["datos"]:
        fallos.append(f"«datos» ha PERDIDO filas: {antes['datos']} → {despues['datos']}")
    return fallos



def _mismo_cliente(ws, filas, cols):
    """⛔ Impide escribir los datos de un cliente en la hoja de OTRO.

    `armar_datos.py` lee todos los volcados que encuentre, y en una pasada de varios
    clientes todos caen en la misma carpeta: sin esto, el segundo cliente arrastraría las
    campañas del primero y acabarían en la hoja equivocada, sin un solo error. Dos
    comprobaciones contra lo que YA hay escrito en la hoja:
      · el nombre de cliente tiene que ser el mismo;
      · y alguna campaña nueva tiene que existir ya ahí (una hoja viva nunca cambia de
        campañas de un día para otro). Si no se parecen en nada, es otro cliente.
    """
    ic, icam = cols["cliente"], cols["campana"]
    clientes_hoja, camp_hoja = set(), set()
    for r in range(2, ws.max_row + 1):
        c, cam = ws.cell(r, ic).value, ws.cell(r, icam).value
        if c: clientes_hoja.add(str(c).strip())
        if cam: camp_hoja.add(str(cam).strip())
    if not clientes_hoja:
        return []                                  # hoja vacía: nada que contrastar
    nuevos = {str(f[ic - 1]).strip() for f in filas}
    camp_nuevas = {str(f[icam - 1]).strip() for f in filas}
    fallos = []
    if nuevos - clientes_hoja:
        fallos.append(f"la hoja es de «{'», «'.join(sorted(clientes_hoja))}» y los datos "
                      f"traen «{'», «'.join(sorted(nuevos))}»")
    if camp_hoja and not (camp_nuevas & camp_hoja):
        fallos.append(f"ninguna de las {len(camp_nuevas)} campañas de los datos está en esta "
                      f"hoja (que tiene {len(camp_hoja)}): parecen de otra cuenta")
    return fallos


def main():
    sheet_id, datos_json = sys.argv[1], sys.argv[2]
    dry = "--dry" in sys.argv
    tmp = os.path.join(tempfile.mkdtemp(), "hoja.xlsx")
    bajar(sheet_id, tmp)
    wb = load_workbook(tmp)                     # con fórmulas: NO data_only
    if "datos" not in wb.sheetnames:
        sys.exit("✗ esa hoja no tiene pestaña `datos`")
    ws = wb["datos"]

    cab = [ws.cell(1, c).value for c in range(1, len(DATOS) + 1)]
    if cab != DATOS:
        sys.exit(f"✗ el orden de columnas de `datos` no es el esperado.\n"
                 f"  hoja:   {cab}\n  script: {DATOS}")

    filas_nuevas = filas_del_json(datos_json)
    cols = {n: i + 1 for i, n in enumerate(DATOS)}
    fallos = _mismo_cliente(ws, filas_nuevas, cols)
    if fallos:
        sys.exit("✗ NO se escribe nada: estos datos no son de esta hoja.\n"
                 + "".join(f"   · {f}\n" for f in fallos))

    # índice de lo que ya hay, por `id`
    donde = {}
    ultima = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v: donde[str(v)] = r; ultima = r

    nuevas = corregidas = iguales = 0
    cambios = []
    for fila in filas_nuevas:
        r = donde.get(fila[0])
        if r:
            antes = [ws.cell(r, c).value for c in range(1, len(DATOS) + 1)]
            # Se compara sin la marca de tiempo (última columna), que siempre cambia. Y se
            # compara NUMÉRICAMENTE: lo que vuelve de la hoja es float («1855.0») y lo nuevo
            # int («1855»), así que compararlo como texto marcaba TODAS las filas como
            # corregidas y escondía cuáles había reatribuido Meta de verdad.
            if not _distintas(antes[:-1], fila[:-1]):
                iguales += 1; continue
            corregidas += 1
            cambios.append((fila[0], _detalle(antes, fila)))
        else:
            ultima += 1; r = ultima; donde[fila[0]] = r; nuevas += 1
        if dry: continue
        for c, (v, fmt) in enumerate(zip(fila, FMTS), start=1):
            cel = ws.cell(r, c); cel.value = v
            if fmt: cel.number_format = fmt

    print(f"  filas nuevas: {nuevas} · corregidas: {corregidas} · sin cambios: {iguales}")
    for id_, det in cambios[:12]:
        print(f"    · {id_}  {det}")
    if len(cambios) > 12: print(f"    … y {len(cambios)-12} más")
    if dry:
        print("  (--dry: no se ha subido nada)"); return
    if not nuevas and not corregidas:
        print("  ya estaba al día: no se sube nada"); return

    ws.auto_filter.ref = f"A1:O{max(ultima, 2)}"
    antes_inv = _inventario(load_workbook(tmp))
    wb.save(tmp)
    fallos = _comprobar(antes_inv, _inventario(load_workbook(tmp)), nuevas)
    if fallos:
        sys.exit("✗ NO se sube nada: el libro ha perdido algo al guardarlo.\n"
                 + "".join(f"   · {f}\n" for f in fallos))
    cuerpo, cabs = _multipart({"mimeType": HOJA}, tmp, XLSX)
    api(f"https://www.googleapis.com/upload/drive/v3/files/{sheet_id}"
        f"?uploadType=multipart&supportsAllDrives=true",
        data=cuerpo, headers=cabs, metodo="PATCH")
    print(f"  ✓ subida al mismo ID · https://docs.google.com/spreadsheets/d/{sheet_id}/edit")


if __name__ == "__main__":
    main()
